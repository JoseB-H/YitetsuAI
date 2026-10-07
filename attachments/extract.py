"""Bounded parsers. Content is data: no macros, formulas, scripts or links execute."""

import csv
from decimal import Decimal, InvalidOperation
import io
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
from html.parser import HTMLParser
import warnings
import zipfile

MAX_CHARACTERS = 250_000
MAX_UNITS = 2000
MAX_ROWS = 5000
MAX_MEDIA_SECONDS = 600


class ExtractionError(ValueError):
    pass


class Result:
    def __init__(self, kind):
        self.data = {
            "kind": kind,
            "units": [],
            "visuals": [],
            "metadata": {},
            "warnings": [],
        }
        self.characters = 0

    def add(self, locator, text):
        text = str(text).strip()
        if not text:
            return
        remaining = MAX_CHARACTERS - self.characters
        if remaining <= 0 or len(self.data["units"]) >= MAX_UNITS:
            self.warn(
                "Extracción parcial: se alcanzó el límite de contenido. Divide el archivo para analizarlo completo."
            )
            return
        if len(text) > remaining:
            self.warn("Extracción parcial: se truncó contenido por el límite de texto.")
        text = text[:remaining]
        self.data["units"].append({"locator": str(locator)[:500], "text": text})
        self.characters += len(text)

    def warn(self, text):
        if text not in self.data["warnings"]:
            self.data["warnings"].append(text)


def decode(path):
    raw = path.read_bytes()
    try:
        text = raw.decode(
            "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        )
    except UnicodeError as exc:
        raise ExtractionError(
            "Texto no UTF-8/UTF-16. Convierte la codificación antes de adjuntarlo."
        ) from exc
    if "\x00" in text:
        raise ExtractionError("El archivo contiene datos binarios, no texto válido")
    return text


def office_container(path, required):
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 2000
                or sum(item.file_size for item in entries) > 50 * 1024 * 1024
            ):
                raise ExtractionError("Documento comprimido demasiado grande")
            if required not in archive.namelist():
                raise ExtractionError(
                    "El contenido no corresponde a la extensión del documento"
                )
            for item in entries:
                if item.filename.startswith("/") or ".." in Path(item.filename).parts:
                    raise ExtractionError("Contenedor de documento inválido")
                if item.filename.lower().endswith(".xml"):
                    content = archive.read(item)
                    if (
                        b"<!DOCTYPE" in content.upper()
                        or b"<!ENTITY" in content.upper()
                    ):
                        raise ExtractionError("Entidades XML no admitidas")
    except zipfile.BadZipFile as exc:
        raise ExtractionError(
            "Documento Office corrupto o con extensión incorrecta"
        ) from exc


def split_text(result, locator, text, size=1400):
    for start in range(0, len(text), size):
        result.add(
            f"{locator}, fragmento {start // size + 1}", text[start : start + size]
        )
        if result.characters >= MAX_CHARACTERS:
            break


def text_file(path, result):
    split_text(result, "texto", decode(path))


def table_summary(rows, columns):
    numeric = {}
    for index, column in enumerate(columns):
        values = []
        for row in rows:
            if index >= len(row) or row[index] in (None, ""):
                continue
            value = row[index]
            # No locale guessing: commas and currency markers remain literal cell data.
            if isinstance(value, bool) or not re.fullmatch(
                r"-?\d+(?:\.\d+)?", str(value)
            ):
                values = []
                break
            try:
                number = Decimal(str(value))
                if number.is_finite():
                    values.append(number)
            except InvalidOperation:
                values = []
                break
        if values:
            numeric[str(column)] = {
                "count": len(values),
                "sum": str(sum(values)),
                "min": str(min(values)),
                "max": str(max(values)),
            }
    return {"rows_included": len(rows), "columns": columns, "numeric_columns": numeric}


def add_rows(result, name, columns, rows):
    for number, row in enumerate(rows, 2):
        result.add(
            f"{name}, fila {number}",
            json.dumps(
                {"columns": columns, "values": row}, ensure_ascii=False, default=str
            ),
        )
    summary = table_summary(rows, columns)
    result.add(
        f"{name}, resumen de filas analizadas",
        json.dumps(summary, ensure_ascii=False, default=str),
    )
    return summary


def csv_file(path, result):
    text = decode(path)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if path.suffix == ".tsv" else csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    columns = [value[:200] for value in next(reader, [])]
    if len(columns) > 100:
        raise ExtractionError("Máximo 100 columnas por tabla")
    rows = []
    for row in reader:
        if len(rows) >= MAX_ROWS:
            result.warn("Tabla parcial: máximo 5000 filas analizadas")
            break
        if len(row) > 100:
            raise ExtractionError("Máximo 100 columnas por tabla")
        rows.append(row)
    result.data["metadata"] = add_rows(result, "tabla", columns, rows)
    result.data["metadata"]["delimiter"] = dialect.delimiter


def json_file(path, result):
    try:
        value = json.loads(
            decode(path),
            parse_constant=lambda value: (_ for _ in ()).throw(
                ValueError("Número no finito")
            ),
        )
    except (ValueError, RecursionError) as exc:
        raise ExtractionError("JSON inválido o demasiado anidado") from exc

    def walk(value, location, depth):
        if depth > 40:
            raise ExtractionError("JSON demasiado anidado")
        if (
            result.characters >= MAX_CHARACTERS
            or len(result.data["units"]) >= MAX_UNITS
        ):
            result.warn("JSON parcial: límite de contenido alcanzado")
            return
        if isinstance(value, dict):
            if not value:
                result.add(location, "{}")
            for key, item in value.items():
                walk(
                    item,
                    location + "[" + json.dumps(key, ensure_ascii=False) + "]",
                    depth + 1,
                )
        elif isinstance(value, list):
            if not value:
                result.add(location, "[]")
            for index, item in enumerate(value):
                walk(item, f"{location}[{index}]", depth + 1)
        else:
            result.add(location, json.dumps(value, ensure_ascii=False))

    walk(value, "$", 0)
    result.data["metadata"] = {"root_type": type(value).__name__}


class VisibleHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.hidden = 0
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.hidden += 1

    def handle_endtag(self, tag):
        if tag in {"script", "style"}:
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def markup_file(path, result):
    source = decode(path)
    if result.data["kind"] == "xml":
        if "<!DOCTYPE" in source.upper() or "<!ENTITY" in source.upper():
            raise ExtractionError("Entidades XML no admitidas")
        from defusedxml.ElementTree import fromstring

        try:
            split_text(result, "XML", " ".join(fromstring(source).itertext()))
        except Exception as exc:
            raise ExtractionError("XML inválido") from exc
    else:
        parser = VisibleHTML()
        parser.feed(source)
        split_text(result, "HTML visible", " ".join(parser.parts))
    result.warn("No se ejecutaron scripts ni se siguieron enlaces externos.")


def ocr(image_path):
    if not shutil.which("tesseract"):
        return "", "OCR no disponible: falta Tesseract"
    languages = subprocess.run(
        ["tesseract", "--list-langs"], capture_output=True, text=True, timeout=10
    ).stdout.splitlines()
    language = "spa+eng" if "spa" in languages else "eng"
    response = subprocess.run(
        ["tesseract", str(image_path), "stdout", "-l", language],
        capture_output=True,
        text=True,
        timeout=45,
    )
    return (
        (response.stdout.strip(), None)
        if response.returncode == 0
        else ("", "No se pudo extraer texto de la imagen")
    )


def visual(path, result, locator):
    from PIL import Image, ImageOps

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(path) as source:
                if source.width * source.height > 25_000_000:
                    raise ExtractionError(
                        "Imagen demasiado grande: máximo 25 megapíxeles"
                    )
                image = ImageOps.exif_transpose(source).convert("RGB")
                image.thumbnail((1536, 1536))
                target = path.parent / f"visual-{len(result.data['visuals'])}.jpg"
                image.save(target, quality=85)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(
            "Imagen corrupta, demasiado grande o no compatible"
        ) from exc
    result.data["visuals"].append({"file": target.name, "locator": locator})
    content, warning = ocr(target)
    if content:
        result.add(locator + ", OCR", content)
    if warning:
        result.warn(warning)
    if path.suffix.lower() == ".gif":
        result.warn("GIF: solo se analiza el primer fotograma.")
    result.warn(
        "OCR y análisis visual pueden equivocarse; verifica cifras y detalles importantes."
    )


def pdf_file(path, result):
    if not path.read_bytes()[:1024].lstrip().startswith(b"%PDF-"):
        raise ExtractionError("El archivo no es un PDF válido")
    from pypdf import PdfReader

    reader = PdfReader(path)
    if reader.is_encrypted:
        raise ExtractionError("PDF cifrado: adjunta una copia sin contraseña")
    if len(reader.pages) > 200:
        raise ExtractionError("Máximo 200 páginas por PDF")
    result.data["metadata"] = {"pages": len(reader.pages)}
    scanned = 0
    for number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if text.strip():
            split_text(result, f"página {number}", text)
        if (
            (not text.strip() or bool(page.images))
            and scanned < 4
            and shutil.which("pdftoppm")
        ):
            prefix = path.parent / f"pdf-page-{number}"
            response = subprocess.run(
                [
                    "pdftoppm",
                    "-f",
                    str(number),
                    "-l",
                    str(number),
                    "-singlefile",
                    "-scale-to",
                    "1536",
                    "-png",
                    str(path),
                    str(prefix),
                ],
                capture_output=True,
                timeout=30,
            )
            if response.returncode == 0:
                visual(prefix.with_suffix(".png"), result, f"página {number}")
                scanned += 1
            else:
                result.warn(f"No se pudo rasterizar la página {number}")
        elif not text.strip() or bool(page.images):
            result.warn(
                f"Página {number}: contenido visual omitido por límite de cuatro páginas o herramienta ausente"
            )


def word_file(path, result):
    office_container(path, "word/document.xml")
    from docx import Document

    document = Document(path)
    for number, paragraph in enumerate(document.paragraphs, 1):
        split_text(result, f"párrafo {number}", paragraph.text)
    for number, table in enumerate(document.tables, 1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        if rows:
            add_rows(result, f"tabla {number}", rows[0], rows[1 : MAX_ROWS + 1])
    with zipfile.ZipFile(path) as archive:
        images = [
            name
            for name in archive.namelist()
            if name.startswith("word/media/")
            and Path(name).suffix.lower() in {".png", ".jpg", ".jpeg", ".gif", ".webp"}
        ]
        for number, name in enumerate(images[:4], 1):
            target = path.parent / f"word-image-{number}{Path(name).suffix.lower()}"
            target.write_bytes(archive.read(name))
            visual(target, result, f"imagen incrustada {number}")
        if len(images) > 4:
            result.warn(
                "Documento parcial: máximo cuatro imágenes incrustadas analizadas."
            )
    result.warn(
        "Se extraen párrafos, tablas e imágenes compatibles; anotaciones y formatos gráficos vectoriales pueden omitirse."
    )


def spreadsheet_file(path, result):
    office_container(path, "xl/workbook.xml")
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        summaries = []
        if len(workbook.worksheets) > 30:
            raise ExtractionError("Máximo 30 hojas por Excel")
        for sheet in workbook.worksheets:
            if sheet.max_column > 100:
                raise ExtractionError("Máximo 100 columnas por hoja")
            iterator = sheet.iter_rows(values_only=True)
            columns = [
                str(value) if value is not None else f"columna {i + 1}"
                for i, value in enumerate(next(iterator, []))
            ]
            rows = []
            for row in iterator:
                if len(rows) >= MAX_ROWS:
                    result.warn(f"Hoja {sheet.title} parcial: máximo 5000 filas")
                    break
                rows.append(list(row))
            summaries.append(
                {
                    "sheet": sheet.title,
                    **add_rows(result, f"hoja {sheet.title}", columns, rows),
                }
            )
        result.data["metadata"] = {"sheets": summaries}
    finally:
        workbook.close()
    result.warn(
        "No se ejecutan fórmulas ni macros; las fórmulas solo aportan su último valor guardado, si existe."
    )


def probe(path):
    if not shutil.which("ffprobe") or not shutil.which("ffmpeg"):
        raise ExtractionError("Procesamiento multimedia no disponible: falta FFmpeg")
    response = subprocess.run(
        [
            "ffprobe",
            "-v",
            "error",
            "-protocol_whitelist",
            "file,pipe",
            "-show_format",
            "-show_streams",
            "-of",
            "json",
            str(path),
        ],
        capture_output=True,
        text=True,
        timeout=20,
    )
    if response.returncode:
        raise ExtractionError("Archivo multimedia corrupto o no compatible")
    data = json.loads(response.stdout)
    duration = float(data.get("format", {}).get("duration", 0))
    if not 0 < duration <= MAX_MEDIA_SECONDS:
        raise ExtractionError(
            "Audio/video debe tener duración conocida de hasta 10 minutos"
        )
    return duration, data.get("streams", [])


def transcribe(path, result):
    from faster_whisper import WhisperModel

    try:
        model = WhisperModel(
            os.getenv("WHISPER_MODEL", "tiny"),
            device="cpu",
            compute_type="int8",
            cpu_threads=2,
            num_workers=1,
            download_root=os.getenv("WHISPER_CACHE_DIR", "data/whisper"),
        )
    except Exception as exc:
        raise ExtractionError(
            "No se pudo cargar Whisper. Descarga el modelo configurado y habilita acceso a Hugging Face o usa WHISPER_MODEL con una ruta local."
        ) from exc
    segments, info = model.transcribe(str(path), beam_size=1, vad_filter=True)
    for segment in segments:
        result.add(f"audio {segment.start:.1f}–{segment.end:.1f} s", segment.text)
    result.data["metadata"]["language"] = info.language
    result.warn(
        "Transcripción automática: puede contener errores, especialmente con ruido, nombres propios o voces superpuestas."
    )


def media_file(path, result):
    duration, streams = probe(path)
    result.data["metadata"]["duration_seconds"] = duration
    has_audio = any(stream.get("codec_type") == "audio" for stream in streams)
    has_video = any(stream.get("codec_type") == "video" for stream in streams)
    if result.data["kind"] == "video" and not has_video:
        raise ExtractionError("El archivo no contiene una pista de video")
    if has_audio:
        wave = path.parent / "audio.wav"
        response = subprocess.run(
            [
                "ffmpeg",
                "-nostdin",
                "-v",
                "error",
                "-protocol_whitelist",
                "file,pipe",
                "-threads",
                "2",
                "-i",
                str(path),
                "-vn",
                "-ac",
                "1",
                "-ar",
                "16000",
                "-t",
                str(MAX_MEDIA_SECONDS),
                "-y",
                str(wave),
            ],
            capture_output=True,
            timeout=45,
        )
        if response.returncode:
            raise ExtractionError("No se pudo decodificar el audio")
        transcribe(wave, result)
    elif result.data["kind"] == "audio":
        raise ExtractionError("El archivo no contiene audio")
    if has_video:
        for i in range(4):
            second = duration * (i + 0.5) / 4
            target = path.parent / f"frame-{i}.jpg"
            response = subprocess.run(
                [
                    "ffmpeg",
                    "-nostdin",
                    "-v",
                    "error",
                    "-protocol_whitelist",
                    "file,pipe",
                    "-threads",
                    "2",
                    "-ss",
                    str(second),
                    "-i",
                    str(path),
                    "-frames:v",
                    "1",
                    "-vf",
                    "scale=1024:-2",
                    "-y",
                    str(target),
                ],
                capture_output=True,
                timeout=30,
            )
            if response.returncode == 0 and target.exists():
                visual(target, result, f"video {second:.1f} s")
            else:
                result.warn(f"No se pudo extraer el fotograma {second:.1f} s")
        result.warn(
            "Video analizado por cuatro fotogramas y audio, si existe; pueden omitirse acciones entre muestras."
        )


def extract(path, kind):
    result = Result(kind)
    handlers = {
        "text": text_file,
        "table": csv_file,
        "structured": json_file,
        "html": markup_file,
        "xml": markup_file,
        "pdf": pdf_file,
        "word": word_file,
        "spreadsheet": spreadsheet_file,
        "audio": media_file,
        "video": media_file,
    }
    try:
        if kind == "image":
            visual(path, result, "imagen")
        else:
            handlers[kind](path, result)
        if not result.data["units"] and not result.data["visuals"]:
            raise ExtractionError("No se encontró contenido interpretable")
        return result.data
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(
            "No se pudo procesar este archivo. Revisa el formato y la integridad."
        ) from exc
