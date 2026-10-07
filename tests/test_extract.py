import asyncio
import io
import json
import os
from pathlib import Path
import subprocess
import tempfile
import unittest
from unittest.mock import patch
from types import SimpleNamespace

from attachments.extract import extract, ExtractionError
from attachments.processing import run_worker
from attachments.storage import safe_filename
from llm import ChatEngine, ModelUnavailable
import httpx


class ExtractionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def file(self, name, content):
        path = self.root / name
        path.write_bytes(content)
        return path

    def test_word_excel_and_json_preserve_structure(self):
        from docx import Document
        from openpyxl import Workbook

        document = Document()
        document.add_paragraph("Aurora launches in October.")
        table = document.add_table(rows=2, cols=2)
        for cell, value in zip(
            [c for row in table.rows for c in row.cells],
            ["Team", "Budget", "Aurora", "30"],
        ):
            cell.text = value
        path = self.root / "report.docx"
        document.save(path)
        word = run_worker(path, "word")
        self.assertIn("October", json.dumps(word))
        self.assertIn("tabla 1", json.dumps(word))
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Expenses"
        sheet.append(["Team", "Cost"])
        sheet.append(["Aurora", 12])
        sheet.append(["Boreal", 18])
        path = self.root / "budget.xlsx"
        workbook.save(path)
        excel = run_worker(path, "spreadsheet")
        self.assertEqual(
            excel["metadata"]["sheets"][0]["numeric_columns"]["Cost"]["sum"], "30"
        )
        data = run_worker(
            self.file(
                "data.json", b'{"team":{"name":"Aurora","active":true},"costs":[12,18]}'
            ),
            "structured",
        )
        self.assertTrue(
            any(
                unit["locator"] == '$["team"]["active"]' and unit["text"] == "true"
                for unit in data["units"]
            )
        )

    def test_pdf_text_and_scanned_page_and_image_ocr(self):
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, NameObject, DictionaryObject
        from PIL import Image, ImageDraw, ImageFont

        writer = PdfWriter()
        page = writer.add_blank_page(width=400, height=200)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): writer._add_object(font)}
                )
            }
        )
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 16 Tf 20 100 Td (Aurora budget 30 dollars) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        path = self.root / "report.pdf"
        writer.write(path)
        data = run_worker(path, "pdf")
        self.assertIn("Aurora budget 30", data["units"][0]["text"])
        self.assertIn("página 1", data["units"][0]["locator"])
        image = Image.new("RGB", (900, 300), "white")
        ImageDraw.Draw(image).text(
            (30, 80),
            "AURORA BUDGET 30",
            font=ImageFont.truetype(
                "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 48
            ),
            fill="black",
        )
        path = self.root / "image.png"
        image.save(path)
        data = run_worker(path, "image")
        self.assertTrue(data["visuals"])
        self.assertIn("AURORA", str(data["units"]))
        pdf = self.root / "scan.pdf"
        image.save(pdf, "PDF")
        data = run_worker(pdf, "pdf")
        self.assertEqual(len(data["visuals"]), 1)
        self.assertIn("AURORA", str(data["units"]))

    def test_markup_does_not_execute_scripts_or_expand_entities(self):
        data = run_worker(
            self.file(
                "report.html",
                b"<h1>Aurora</h1><script>secret-script</script><p>Budget 30</p>",
            ),
            "html",
        )
        self.assertNotIn("secret-script", str(data["units"]))
        with self.assertRaises(ExtractionError):
            run_worker(
                self.file(
                    "bad.xml",
                    b'<!DOCTYPE x [<!ENTITY boom SYSTEM "file:///etc/passwd">]><x>&boom;</x>',
                ),
                "xml",
            )

    def test_video_decoding_samples_frames_without_audio(self):
        path = self.root / "clip.mp4"
        subprocess.run(
            [
                "ffmpeg",
                "-nostdin",
                "-v",
                "error",
                "-f",
                "lavfi",
                "-i",
                "color=c=red:s=320x240:d=2",
                "-c:v",
                "mpeg4",
                "-y",
                str(path),
            ],
            check=True,
        )
        data = run_worker(path, "video")
        self.assertEqual(len(data["visuals"]), 4)
        self.assertAlmostEqual(data["metadata"]["duration_seconds"], 2, places=1)
        self.assertTrue(any("fotogramas" in warning for warning in data["warnings"]))

    def test_audio_decoding_and_time_locators_with_stubbed_transcriber(self):
        # Real media decoder; only speech inference is stubbed. Live speech needs model download.
        path = self.root / "speech.wav"
        subprocess.run(
            [
                "ffmpeg",
                "-nostdin",
                "-v",
                "error",
                "-f",
                "lavfi",
                "-i",
                "sine=frequency=440:duration=1",
                "-y",
                str(path),
            ],
            check=True,
        )

        class SpeechStub:
            def __init__(self, *args, **kwargs):
                pass

            def transcribe(self, *args, **kwargs):
                return [
                    SimpleNamespace(start=0, end=1, text="Aurora budget 30")
                ], SimpleNamespace(language="en")

        with patch("faster_whisper.WhisperModel", SpeechStub):
            data = extract(path, "audio")
        self.assertEqual(data["units"][0]["locator"], "audio 0.0–1.0 s")
        self.assertEqual(data["metadata"]["language"], "en")
        self.assertTrue((self.root / "audio.wav").exists())

    def test_office_zip_bombs_and_binary_text_are_rejected(self):
        import zipfile

        path = self.root / "bomb.docx"
        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", b" " * (51 * 1024 * 1024))
        with self.assertRaises(ExtractionError):
            run_worker(path, "word")
        with self.assertRaises(ExtractionError):
            run_worker(self.file("binary.txt", b"hello\x00world"), "text")


class ModelTransportTests(unittest.IsolatedAsyncioTestCase):
    async def test_real_chat_engine_supplies_structured_file_context_and_history(self):
        captured = []

        def respond(request):
            captured.append(json.loads(request.content))
            return httpx.Response(
                200,
                json={
                    "message": {
                        "content": "Aurora y Boreal suman 30 (ventas.csv, resumen)."
                    }
                },
            )

        client = httpx.AsyncClient(
            transport=httpx.MockTransport(respond), base_url="http://ollama"
        )
        model = ChatEngine(client)
        item = SimpleNamespace(
            id="00000000-0000-0000-0000-000000000001",
            filename="ventas.csv",
            kind="table",
            extraction={
                "units": [{"locator": "tabla, resumen", "text": "sum=30"}],
                "metadata": {},
                "visuals": [],
                "warnings": [],
            },
        )
        result = await model.answer(
            "Compara ingresos",
            [SimpleNamespace(role="user", content="Recuerda Aurora")],
            [item],
        )
        self.assertIn("ventas.csv", json.dumps(captured[0]))
        self.assertIn("Recuerda Aurora", json.dumps(captured[0]))
        self.assertEqual(result["sources"][0]["filename"], "ventas.csv")
        await model.close()

    async def test_provider_error_never_becomes_a_successful_answer(self):
        client = httpx.AsyncClient(
            transport=httpx.MockTransport(
                lambda request: httpx.Response(404, json={"error": "model missing"})
            ),
            base_url="http://ollama",
        )
        model = ChatEngine(client)
        with self.assertRaises(ModelUnavailable):
            await model.answer("Hello", [], [])
        await model.close()
